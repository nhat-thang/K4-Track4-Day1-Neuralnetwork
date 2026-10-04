"""results_table.py — bản hoàn thiện (Lab Day 1).

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path

COLUMNS = [
    "exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay", "batch", "epochs", "hidden",
    "dropout", "clip_norm", "precision", "init", "seed", "step0_loss", "best_val_loss", "best_epoch",
    "final_train_loss", "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB",
    "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes",
]
FORMULA_COLUMNS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
OPT_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}
GROUP_ALIASES = {"lr_search": "hparam"}   # lần chạy tìm lr của Part 2 thuộc chủ đề hyper-parameter


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    out_dir = Path(results_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    cfg = {k: (list(v) if isinstance(v, tuple) else v) for k, v in result["cfg"].items()}
    payload = {"cfg": cfg, "history": result["history"], "summary": result["summary"]}
    path = out_dir / f"{cfg['exp_id']}.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    results = []
    for path in sorted(Path(results_dir).glob("*.json")):
        with open(path, encoding="utf-8") as f:
            r = json.load(f)
        if not {"cfg", "history", "summary"} <= set(r):   # bỏ qua file JSON không phải kết quả một lần chạy
            continue
        assert r["cfg"]["exp_id"] == path.stem, f"{path.name}: exp_id trong file khác tên file"
        results.append(r)
    return sorted(results, key=lambda r: r["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng.

    Định dạng theo từ điển cột của mẫu: loss CE|MSE, optimizer SGD|SGD+momentum|Adam|AdamW,
    hidden "256-128", clip_norm "none" hoặc số, diverged Y|N. Giá trị thiếu (None) để trống."""
    cfg, summ = result["cfg"], result["summary"]
    row = {
        "exp_id": cfg["exp_id"],
        "group": GROUP_ALIASES.get(cfg["group"], cfg["group"]),
        "description": cfg["description"],
        "loss": cfg["loss"].upper(),
        "optimizer": OPT_NAMES[cfg["optimizer"]],
        "lr": cfg["lr"],
        "weight_decay": cfg["weight_decay"],
        "batch": cfg["batch"],
        "epochs": cfg["epochs"],
        "hidden": "-".join(str(h) for h in cfg["hidden"]),
        "dropout": cfg["dropout"],
        "clip_norm": "none" if cfg["clip_norm"] is None else cfg["clip_norm"],
        "precision": cfg["precision"],
        "init": cfg["init"],
        "seed": cfg["seed"],
    }
    for k in ("step0_loss", "best_val_loss", "best_epoch", "final_train_loss", "final_val_loss",
              "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB"):
        row[k] = summ.get(k)
    row["diverged"] = "Y" if summ.get("diverged") else "N"
    if eval_scores is not None:
        row["eval_acc"] = eval_scores["eval_acc"]
        row["eval_macro_f1"] = eval_scores["eval_macro_f1"]
    row["figure_file"] = f"figures/{cfg['exp_id']}.png"
    extra = " | ".join(t for t in (cfg.get("notes", ""), notes) if t)
    row["notes"] = extra
    assert set(row) <= set(COLUMNS), set(row) - set(COLUMNS)
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    seed_ids     : exp_id các lần chạy baseline khác seed, ghi vào cột A (dòng 2..6) của sheet "Seeds".
    summary_notes: {group: nhận xét ngắn} ghi vào cột H của sheet "Summary" (hàng có cột A = group).

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    import openpyxl

    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]
    col_of = {c.value: c.column for c in ws[1] if c.value}
    missing = [c for c in COLUMNS if c not in col_of]
    assert not missing, f"mẫu thiếu cột: {missing}"
    n_rows = ws.max_row - 1
    assert len(rows) <= n_rows, f"mẫu chỉ có {n_rows} dòng, cần {len(rows)}"
    ids = [r["exp_id"] for r in rows]
    assert len(set(ids)) == len(ids), "exp_id bị trùng"

    for col in COLUMNS:                          # xoá dòng mẫu có sẵn (dòng baseline điền sẵn) trước khi ghi
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col_of[col]).value = None
    for i, row in enumerate(rows):
        for key, val in row.items():
            assert key not in FORMULA_COLUMNS, f"không ghi đè cột công thức {key}"
            ws.cell(row=i + 2, column=col_of[key]).value = val

    if seed_ids is not None:
        seeds = wb["Seeds"]
        assert len(seed_ids) <= 5
        for r in range(2, 7):
            seeds.cell(row=r, column=1).value = seed_ids[r - 2] if r - 2 < len(seed_ids) else None
    if summary_notes:
        summ = wb["Summary"]
        for r in range(2, 12):
            g = summ.cell(row=r, column=1).value
            if g in summary_notes:
                summ.cell(row=r, column=8).value = summary_notes[g]
    wb.save(out_path)
